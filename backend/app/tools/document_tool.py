import os
import logging
from typing import Dict, Any, List

import docx
from policy import permissions

logger = logging.getLogger("tools.document_tool")

def create_docx(filepath: str, title: str, paragraphs: List[str]) -> Dict[str, Any]:
    """
    Creates a .docx file with a title and paragraph content.
    """
    if not filepath:
        return {"success": False, "error": "No file path provided."}
        
    # Permission check
    if not permissions.check_file_write_permission(filepath):
        return {"success": False, "error": f"Write permission denied for path: {filepath}"}
        
    try:
        norm_path = os.path.normpath(os.path.abspath(filepath))
        # Ensure parent folder exists
        parent_dir = os.path.dirname(norm_path)
        if parent_dir and not os.path.exists(parent_dir):
            os.makedirs(parent_dir, exist_ok=True)
            
        doc = docx.Document()
        
        # Add Title
        if title:
            doc.add_heading(title, level=0)
            
        # Add Paragraphs
        for p in paragraphs:
            doc.add_paragraph(p)
            
        doc.save(norm_path)
        logger.info(f"Created docx document: {norm_path}")
        return {"success": True, "message": f"Successfully created document: {norm_path}"}
    except Exception as e:
        logger.error(f"Failed to create docx {filepath}: {e}")
        return {"success": False, "error": str(e)}

def read_text(filepath: str) -> Dict[str, Any]:
    """
    Reads text content from .txt or .docx files.
    """
    if not filepath:
        return {"success": False, "error": "No file path provided."}
        
    # Permission check
    if not permissions.check_file_read_permission(filepath):
        return {"success": False, "error": f"Read permission denied for path: {filepath}"}
        
    if not os.path.exists(filepath):
        return {"success": False, "error": f"File not found: {filepath}"}
        
    _, ext = os.path.splitext(filepath)
    ext = ext.lower()
    
    try:
        norm_path = os.path.normpath(os.path.abspath(filepath))
        text_content = ""
        
        if ext == ".txt":
            with open(norm_path, "r", encoding="utf-8", errors="ignore") as f:
                text_content = f.read()
        elif ext == ".docx":
            doc = docx.Document(norm_path)
            paragraphs_text = [p.text for p in doc.paragraphs]
            # Include table texts
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs_text.append(cell.text)
            text_content = "\n".join(paragraphs_text)
        else:
            return {"success": False, "error": f"Unsupported extension for direct text reading: {ext}"}
            
        return {"success": True, "content": text_content, "length": len(text_content)}
    except Exception as e:
        logger.error(f"Failed to read text from {filepath}: {e}")
        return {"success": False, "error": str(e)}

def summarize_text(filepath: str) -> Dict[str, Any]:
    """
    Reads the file and generates a simple text summary.
    """
    res = read_text(filepath)
    if not res["success"]:
        return res
        
    content = res["content"]
    if not content:
        return {"success": True, "summary": "Tệp tin trống."}
        
    # Simple rule-based summarization for Phase 2 local logic
    lines = [line.strip() for line in content.split("\n") if line.strip()]
    if not lines:
        return {"success": True, "summary": "Tệp tin không có nội dung văn bản hợp lệ."}
        
    # Take first 3 lines and total word count
    words = content.split()
    total_words = len(words)
    summary_preview = "\n".join(lines[:3])
    
    summary_text = (
        f"Bản tóm tắt tệp tin:\n"
        f"- Tổng số từ: {total_words}\n"
        f"- Tổng số dòng: {len(lines)}\n"
        f"- Nội dung chính sơ lược:\n{summary_preview}"
    )
    
    return {"success": True, "summary": summary_text}
